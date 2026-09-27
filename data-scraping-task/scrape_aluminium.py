"""Scrape business contacts from Rainbow Pages -> Construction -> Aluminium.

Walks every pagination page of the category, opens each company profile to
read the (Cloudflare-obfuscated) email, removes duplicate companies and writes
a single Excel file.
"""

import logging
import re
import sys
import time
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

CATEGORY_URL = "https://rainbowpages.lk/constructions/aluminium/"
PAGE_URL = CATEGORY_URL + "?page={page}&s=Aluminium&l="
OUTPUT_FILE = Path(__file__).parent / "output" / "rainbowpages_aluminium_contacts.xlsx"
NOT_FOUND = "Not found."
REQUEST_DELAY = 1.0  # seconds between requests, to be polite to the server
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")

log = logging.getLogger("scrape_aluminium")


# --------------------------------------------------------------------------- #
# HTTP
# --------------------------------------------------------------------------- #
def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(HEADERS)
    retry = Retry(total=3, backoff_factor=2, status_forcelist=(429, 500, 502, 503, 504))
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session


_last_request = 0.0


def fetch(session: requests.Session, url: str) -> str:
    global _last_request
    wait = REQUEST_DELAY - (time.monotonic() - _last_request)
    if wait > 0:
        time.sleep(wait)
    resp = session.get(url, timeout=30)
    _last_request = time.monotonic()
    resp.raise_for_status()
    return resp.text


# --------------------------------------------------------------------------- #
# Parsing helpers
# --------------------------------------------------------------------------- #
def clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def normalize_phones(*phone_strings: str) -> str:
    seen, phones = set(), []
    for s in phone_strings:
        for p in re.split(r"[,/;]", s or ""):
            p = clean(p)
            key = re.sub(r"\D", "", p)
            if p and key not in seen:
                seen.add(key)
                phones.append(p)
    return ", ".join(phones)


def decode_cfemail(hex_str: str) -> str:
    """Decode Cloudflare email protection: first byte is the XOR key."""
    data = bytes.fromhex(hex_str)
    return bytes(b ^ data[0] for b in data[1:]).decode("utf-8", errors="replace")


def get_max_page(soup: BeautifulSoup) -> int:
    pages = [int(m) for a in soup.select("a[href*='page=']")
             for m in re.findall(r"page=(\d+)", a["href"])]
    return max(pages, default=1)


def parse_listing_page(html: str, page_no: int) -> list[dict]:
    soup = BeautifulSoup(html, "lxml")
    listings = []
    for body in soup.select(".media-body"):
        link = body.select_one("h4.media-heading a")
        if not link or not link.get("href"):
            continue
        address = phones = ""
        for p in body.find_all("p", recursive=False):
            if p.select_one("i.fa-map-marker"):
                address = clean(p.get_text())
            elif p.select_one("i.fa-phone"):
                phones = clean(p.get_text())
        listings.append({
            "name": clean(link.get_text()),
            "profile_url": urljoin(CATEGORY_URL, link["href"]),
            "address": address,
            "phones": phones,
            "page": page_no,
        })
    return listings


def parse_profile(html: str) -> dict:
    """Read Address / Telephone / Email from the profile's Contact Details block."""
    soup = BeautifulSoup(html, "lxml")
    block = soup.select_one(".contact-details")
    fields: dict[str, object] = {}
    if not block:
        return fields
    for row in block.select(".row"):
        label = row.find("strong")
        cols = row.find_all("div", recursive=False)
        if not label or len(cols) < 2:
            continue
        name, value = clean(label.get_text()).lower(), cols[1]
        if name == "address":
            fields["address"] = clean(value.get_text())
        elif name == "telephone":
            fields["phones"] = clean(value.get_text())
        elif name == "email":
            fields["emails"] = extract_emails(value)
    return fields


def extract_emails(node) -> list[str]:
    """Only emails actually present on the page: decoded cfemail or plain mailto/text."""
    found = []
    for el in node.select("[data-cfemail]"):
        found.append(decode_cfemail(el["data-cfemail"]))
    for a in node.select("a[href^='mailto:']"):
        found.append(a["href"][len("mailto:"):].split("?")[0])
    if not found:
        found.extend(re.findall(r"[\w.%+'-]+@[\w.-]+\.[A-Za-z]{2,}", node.get_text()))
    emails, seen = [], set()
    for e in found:
        e = clean(e).lower()
        if EMAIL_RE.match(e) and e not in seen:
            seen.add(e)
            emails.append(e)
    return emails


# --------------------------------------------------------------------------- #
# Dedupe
# --------------------------------------------------------------------------- #
def norm_name(name: str) -> str:
    n = name.lower().replace("&", " and ")
    n = re.sub(r"\b(private|pvt|pte)\b", "pvt", n)
    n = re.sub(r"\blimited\b", "ltd", n)
    return re.sub(r"[^a-z0-9]", "", n)


def norm_address(address: str) -> str:
    return re.sub(r"[^a-z0-9]", "", address.lower())


def phone_set(phones: str) -> set[str]:
    return {d[-9:] for d in (re.sub(r"\D", "", p) for p in phones.split(",")) if len(d) >= 9}


def same_company(a: dict, b: dict) -> bool:
    """Same business listed twice: identical name with matching address or a shared
    phone, or a near-identical name (spelling variant) with a shared phone."""
    na, nb = norm_name(a["name"]), norm_name(b["name"])
    shared_phone = bool(phone_set(a["phones"]) & phone_set(b["phones"]))
    if na == nb:
        aa, ab = norm_address(a["address"]), norm_address(b["address"])
        return shared_phone or aa == ab or (aa and ab and (aa in ab or ab in aa))
    return shared_phone and SequenceMatcher(None, na, nb).ratio() >= 0.85


def merge_addresses(addresses: list[str]) -> str:
    """Keep distinct branch addresses; drop ones contained in, or a typo of, another."""
    kept: list[str] = []
    for addr in sorted(addresses, key=len, reverse=True):
        na = norm_address(addr)
        if addr and not any(na in norm_address(k)
                            or SequenceMatcher(None, na, norm_address(k)).ratio() >= 0.9
                            for k in kept):
            kept.append(addr)
    return " | ".join(kept)


def dedupe_companies(records: list[dict]) -> list[dict]:
    # Union-find over all pairs, so chains of duplicates collapse into one group.
    parent = list(range(len(records)))

    def root(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i in range(len(records)):
        for j in range(i + 1, len(records)):
            if same_company(records[i], records[j]):
                parent[root(j)] = root(i)

    groups: dict[int, list[dict]] = {}
    for i, r in enumerate(records):
        groups.setdefault(root(i), []).append(r)

    merged = []
    for group in groups.values():
        emails: list[str] = []
        for r in group:
            emails += [e for e in r["emails"] if e not in emails]
        # Prefer the name from a listing that has an email, else the first seen.
        best = next((r for r in group if r["emails"]), group[0])
        merged.append({
            "name": best["name"],
            "address": merge_addresses([r["address"] for r in group]),
            "phones": normalize_phones(*(r["phones"] for r in group)),
            "emails": emails,
        })
        if len(group) > 1:
            log.info("Merged duplicates: %s", " / ".join(r["name"] for r in group))
    return merged


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
def export_xlsx(rows: list[dict], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Aluminium Contacts"
    headers = ["Company Name", "Address", "Phone Numbers", "Email Address"]
    ws.append(headers)
    for r in rows:
        ws.append([r["name"], r["address"], r["phones"] or NOT_FOUND,
                   ", ".join(r["emails"]) or NOT_FOUND])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCD", (42, 50, 40, 38)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(path)
    except PermissionError:
        # File is likely open in Excel; don't throw away a finished scrape.
        path = path.with_name(f"{path.stem}_{time.strftime('%Y%m%d_%H%M%S')}{path.suffix}")
        log.warning("Output file is locked (open in Excel?); saving to %s instead", path.name)
        wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    session = make_session()

    # 1. Walk every listing page.
    first_html = fetch(session, CATEGORY_URL)
    max_page = get_max_page(BeautifulSoup(first_html, "lxml"))
    log.info("Category has %d pages", max_page)

    raw = parse_listing_page(first_html, 1)
    log.info("Page 1/%d: %d listings", max_page, len(raw))
    for page in range(2, max_page + 1):
        items = parse_listing_page(fetch(session, PAGE_URL.format(page=page)), page)
        log.info("Page %d/%d: %d listings", page, max_page, len(items))
        raw.extend(items)

    # 2. Drop repeated profile URLs so no profile is fetched twice.
    by_url: dict[str, dict] = {}
    for item in raw:
        by_url.setdefault(item["profile_url"], item)
    listings = list(by_url.values())
    log.info("Raw listings: %d, unique profiles: %d", len(raw), len(listings))

    # 3. Open each profile for the email (and fuller address/phones).
    records, failed = [], 0
    for i, item in enumerate(listings, 1):
        try:
            prof = parse_profile(fetch(session, item["profile_url"]))
        except requests.RequestException as exc:
            log.warning("Profile failed (%s): %s", item["profile_url"], exc)
            prof, failed = {}, failed + 1
        records.append({
            "name": item["name"],
            "address": prof.get("address") or item["address"],
            "phones": normalize_phones(prof.get("phones") or item["phones"]),
            "emails": prof.get("emails", []),
        })
        if i % 25 == 0 or i == len(listings):
            log.info("Profiles fetched: %d/%d", i, len(listings))

    # 4. Remove duplicate companies and export.
    final = sorted(dedupe_companies(records), key=lambda r: r["name"].lower())
    saved = export_xlsx(final, OUTPUT_FILE)

    with_email = sum(1 for r in final if r["emails"])
    log.info("Pages scraped: %d | raw listings: %d | duplicates removed: %d | final companies: %d",
             max_page, len(raw), len(raw) - len(final), len(final))
    log.info("Emails found: %d | %s: %d | profile fetch failures: %d",
             with_email, NOT_FOUND, len(final) - with_email, failed)
    log.info("Saved %s", saved)
    return 0


if __name__ == "__main__":
    sys.exit(main())
